"""Import of indicators from a file provided by the analyst (txt, csv, xlsx).

Deterministic pattern-based extraction, without going through a model: defang
normalized, classification by shape (hash, public IP, domain, URL, email, path, registry
key), deduplication. The extracted indicators arrive pending validation: the import
avoids manual entry, not the human checkpoint.
"""

from __future__ import annotations

import io
import ipaddress
import re
from urllib.parse import urlsplit

from middleware.errors import ErrorCode, ToolError
from middleware.guardrails.ioc import IocType

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_INDICATORS = 2000

_HASH_RE = re.compile(r"^(?:[a-f0-9]{32}|[a-f0-9]{40}|[a-f0-9]{64}|[a-f0-9]{128})$", re.I)
_DOMAIN_RE = re.compile(r"^(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,24}$", re.I)
_EMAIL_RE = re.compile(r"^[\w.+-]+@[\w-]+(?:\.[\w-]+)*\.[a-z]{2,}$", re.I)
_IDENTIFIER_RE = re.compile(r"^@?[a-z0-9][\w.@/+-]{2,127}$", re.I)
_HAS_LETTER_RE = re.compile(r"[a-z]", re.I)
_WINDOWS_PATH_RE = re.compile(r"^[a-z]:\\[^\s]{2,}", re.I)
_UNIX_PATH_RE = re.compile(
    r"^/(?:tmp|var|usr|etc|home|opt|bin|srv|dev|private|library)/[^\s]+", re.I
)
_REGISTRY_RE = re.compile(r"^(?:HKEY_[A-Z_]+|HKLM|HKCU|HKCR|HKU)\\", re.I)
_SPLIT_RE = re.compile(r"[\s,;\"'<>()\[\]{}|]+")
_WS_RE = re.compile(r"\s")

_FILE_EXTENSIONS = {
    "pdf",
    "txt",
    "csv",
    "xls",
    "xlsx",
    "doc",
    "docx",
    "ppt",
    "pptx",
    "exe",
    "dll",
    "bat",
    "ps1",
    "vbs",
    "js",
    "jar",
    "zip",
    "rar",
    "7z",
    "gz",
    "tar",
    "iso",
    "img",
    "msi",
    "scr",
    "lnk",
    "tmp",
    "log",
    "dat",
    "bin",
    "sys",
    "ini",
    "cfg",
    "conf",
    "yml",
    "yaml",
    "json",
    "xml",
    "html",
    "htm",
    "php",
    "asp",
    "aspx",
    "jsp",
    "py",
    "sh",
    "md",
    "png",
    "jpg",
    "jpeg",
    "gif",
    "bmp",
    "svg",
    "mp3",
    "mp4",
    "avi",
}


def extract_text(data: bytes, filename: str) -> str:
    """Raw text of an import file. xlsx via openpyxl, everything else is decoded as text
    (txt, csv, exported lists)."""

    if len(data) > MAX_FILE_BYTES:
        raise ToolError(
            ErrorCode.SCHEMA_INVALID,
            f"File too large (cap {MAX_FILE_BYTES // (1024 * 1024)} MB).",
        )
    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as error:  # openpyxl raises varied types on a corrupted file
            raise ToolError(ErrorCode.SCHEMA_INVALID, "Unreadable Excel file.") from error
        parts: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                parts.extend(str(cell) for cell in row if cell is not None)
        workbook.close()
        return "\n".join(parts)
    if filename.lower().endswith((".xls", ".ods")):
        raise ToolError(
            ErrorCode.SCHEMA_INVALID,
            "Unsupported format: export to .xlsx, .csv or .txt.",
        )
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def _is_public_ip(value: str) -> bool:
    try:
        address = ipaddress.ip_address(value)
    except ValueError:
        return False
    return not (
        address.is_private
        or address.is_loopback
        or address.is_link_local
        or address.is_reserved
        or address.is_multicast
        or address.is_unspecified
    )


def _classify(token: str, *, allow_other: bool = False) -> tuple[str, IocType] | None:
    """The defang is already normalized at the text level. Detection is done on the
    lowercase form; the kept value keeps its case for what is case-sensitive (paths,
    registry keys, URLs)."""

    raw = token.strip().strip("\"'").rstrip(".,;:")
    if len(raw) < 4:
        return None
    low = raw.lower()
    if _HASH_RE.match(low):
        return low, IocType.HASH
    if _REGISTRY_RE.match(raw):
        return raw, IocType.REGISTRY_KEY
    if _WINDOWS_PATH_RE.match(raw) or _UNIX_PATH_RE.match(raw):
        return raw, IocType.FILE_PATH
    if low.startswith(("http://", "https://")):
        parts = urlsplit(raw)
        if parts.hostname:
            return raw, IocType.URL
        return None
    if _EMAIL_RE.match(low):
        return low, IocType.EMAIL
    if _is_public_ip(raw):
        return raw, IocType.IP
    if _DOMAIN_RE.match(low):
        suffix = low.rsplit(".", 1)[-1]
        if suffix in _FILE_EXTENSIONS:
            return None
        return low, IocType.DOMAIN
    if allow_other and _IDENTIFIER_RE.match(raw) and _HAS_LETTER_RE.search(raw):
        # List mode: a token that matches no network/host category but has the shape of an
        # identifier (package name, mutex, wallet...) is kept as "other" rather than
        # dropped or misclassified. The agent will search for it as a literal.
        return raw, IocType.OTHER
    return None


def extract_indicators(text: str) -> list[tuple[str, IocType]]:
    """Distinct indicators found in a text, capped at MAX_INDICATORS.

    Deterministic and conservative: an ambiguous token (plain file name, private IP, any
    word) is ignored rather than guessed.
    """

    text = text.replace("[.]", ".").replace("(.)", ".").replace("hxxp", "http")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    single = [line for line in lines if line and not _WS_RE.search(line)]
    # "List" file (one indicator per line, mostly): each line is an indicator, with an
    # "other" fallback for out-of-category shapes (package names). Otherwise (prose
    # report), only recognized shapes are kept, without "other", so as not to turn every
    # word into an indicator.
    list_mode = len(lines) >= 3 and len(single) >= 0.6 * len(lines)

    found: dict[str, tuple[str, IocType]] = {}

    def _add(candidate: tuple[str, IocType] | None) -> bool:
        if candidate is None:
            return True
        value, ioc_type = candidate
        key = value.lower()
        if key not in found:
            found[key] = (value, ioc_type)
        return len(found) < MAX_INDICATORS

    if list_mode:
        for line in lines:
            if _WS_RE.search(line):
                # compound line: fall back to token-based extraction (standalone shapes)
                for token in _SPLIT_RE.split(line):
                    token = token.strip().strip(".").strip()
                    if token and not _add(_classify(token)):
                        return list(found.values())
                continue
            token = line.strip().strip(".").strip()
            if token and not _add(_classify(token, allow_other=True)):
                break
        return list(found.values())

    for token in _SPLIT_RE.split(text):
        token = token.strip().strip(".").strip()
        if token and not _add(_classify(token)):
            break
    return list(found.values())
